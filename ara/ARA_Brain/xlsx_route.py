# -*- coding: utf-8 -*-
"""
Módulo de generación de Reporte XLSX - Sistema ARA
Endpoint: /api/reporte/xlsx (GET y POST)

Mismo dataset que pdf_route.py (Dashboard de Rendimiento), pero como libro de
Excel profesional con 4 hojas, encabezados fijos, autofiltro y formato de
tabla — reemplaza el PDF porque el equipo necesita filtrar/ordenar/pivotar
los datos, algo que un PDF no permite.
"""
import io
from datetime import datetime, timedelta

from flask import request, send_file, make_response, jsonify
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pdf_route import (
    obtener_incidencias_rango,
    obtener_resumen_kpis,
    obtener_distribucion_por_modulo,
    obtener_rendimiento_operadores,
    _label_modulo,
)

# =============================================================================
# ESTILO CORPORATIVO (mismos tonos navy del PDF: #1a1a2e / #16213e)
# =============================================================================
_NAVY = 'FF1A1A2E'
_NAVY_2 = 'FF16213E'
_GRIS_CLARO = 'FFF0F0F5'
_BLANCO = 'FFFFFFFF'

_FONT_HEADER = Font(name='Calibri', size=11, bold=True, color=_BLANCO)
_FONT_SUBTITULO = Font(name='Calibri', size=11, bold=True, color='16213E')
_FILL_HEADER = PatternFill('solid', fgColor=_NAVY)
_FILL_HEADER_2 = PatternFill('solid', fgColor=_NAVY_2)
_FILL_ALT = PatternFill('solid', fgColor=_GRIS_CLARO)
_BORDE_FINO = Border(*[Side(style='thin', color='B0B0C0')] * 4)
_ALINEAR_CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
_ALINEAR_IZQ = Alignment(horizontal='left', vertical='center')


def _escribir_tabla(ws, fila_inicio, headers, filas, anchos=None, alterno=True, filtro=False):
    """Escribe una tabla con encabezado navy + filas alternadas + bordes.

    Retorna la fila siguiente a la última escrita (para encadenar bloques).
    """
    for c, h in enumerate(headers, start=1):
        cel = ws.cell(row=fila_inicio, column=c, value=h)
        cel.font = _FONT_HEADER
        cel.fill = _FILL_HEADER
        cel.alignment = _ALINEAR_CENTRO
        cel.border = _BORDE_FINO
    ws.row_dimensions[fila_inicio].height = 22

    for i, fila in enumerate(filas):
        r = fila_inicio + 1 + i
        for c, val in enumerate(fila, start=1):
            cel = ws.cell(row=r, column=c, value=val)
            cel.border = _BORDE_FINO
            cel.alignment = _ALINEAR_CENTRO if c > 1 else _ALINEAR_IZQ
            if alterno and i % 2 == 1:
                cel.fill = _FILL_ALT

    if anchos:
        for c, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w

    if filtro and filas:
        ultima_fila = fila_inicio + len(filas)
        ultima_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{fila_inicio}:{ultima_col}{ultima_fila}"

    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1).coordinate

    return fila_inicio + len(filas) + 1


def _hoja_portada(wb, fecha_inicio, fecha_fin, kpis):
    ws = wb.active
    ws.title = 'Resumen'
    ws['B2'] = 'REPORTE DE RENDIMIENTO — SISTEMA ARA'
    ws['B2'].font = Font(name='Calibri', size=18, bold=True, color=_NAVY.replace('FF', ''))
    ws['B3'] = f'Período: {fecha_inicio} al {fecha_fin}'
    ws['B3'].font = _FONT_SUBTITULO
    ws['B4'] = f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws['B4'].font = Font(italic=True, color='64748B')

    kpi_items = [
        ('Total Operaciones', kpis.get('total_operaciones') or 0),
        ('Total Renglones', kpis.get('total_renglones') or 0),
        ('Total Puntos', round(kpis.get('total_puntos') or 0, 2)),
        ('Usuarios Activos', kpis.get('usuarios_activos') or 0),
    ]
    fila = 6
    for etiqueta, valor in kpi_items:
        ws.cell(row=fila, column=2, value=etiqueta).font = Font(bold=True, color='475569')
        c = ws.cell(row=fila, column=3, value=valor)
        c.font = Font(size=14, bold=True, color=_NAVY.replace('FF', ''))
        c.alignment = _ALINEAR_CENTRO
        fila += 1

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 18
    ws.sheet_view.showGridLines = False


def _hoja_distribucion(wb, fecha_inicio, fecha_fin, distribucion_modulo):
    ws = wb.create_sheet('Distribución por Módulo')
    ws['A1'] = f'Distribución de operaciones por módulo ({fecha_inicio} a {fecha_fin})'
    ws['A1'].font = _FONT_SUBTITULO

    total_renglones = sum((m.get('renglones') or 0) for m in distribucion_modulo) or 1
    headers = ['Módulo', 'Operaciones', 'Renglones', '% del Volumen', 'Puntos', 'Operadores']
    filas = []
    for m in distribucion_modulo:
        renglones = m.get('renglones') or 0
        filas.append([
            _label_modulo(m.get('modulo')),
            m.get('operaciones') or 0,
            renglones,
            round(renglones / total_renglones * 100, 1),
            round(m.get('puntos') or 0, 2),
            m.get('operadores') or 0,
        ])
    _escribir_tabla(ws, 3, headers, filas, anchos=[26, 14, 12, 14, 12, 12], filtro=True)
    ws.sheet_view.showGridLines = False


def _hoja_operadores(wb, fecha_inicio, fecha_fin, rendimiento_operadores):
    ws = wb.create_sheet('Rendimiento por Operador')
    ws['A1'] = f'Rendimiento por operador ({fecha_inicio} a {fecha_fin})'
    ws['A1'].font = _FONT_SUBTITULO

    total_ops = sum((o.get('total_operaciones') or 0) for o in rendimiento_operadores) or 1
    modulos_clave = ['picking', 'chequeo', 'inventario', 'embalaje']
    headers = [
        'Operador', '% Operaciones', 'Renglones', 'Puntos Totales',
        'Pts Picking', 'Pts Chequeo', 'Pts Inventario', 'Pts Embalaje',
        'Punto Fuerte', 'Área de Mejora',
    ]
    filas = []
    for op in rendimiento_operadores:
        pct_ops = round((op.get('total_operaciones') or 0) / total_ops * 100, 1)
        puntos_por_modulo = {m: (op.get(f'puntos_{m}') or 0) for m in modulos_clave}
        renglones_por_modulo = {m: (op.get(f'renglones_{m}') or 0) for m in modulos_clave}
        modulos_activos = {m: p for m, p in puntos_por_modulo.items() if renglones_por_modulo[m] > 0}

        if modulos_activos:
            m_fuerte = max(modulos_activos, key=modulos_activos.get)
            fuerte_txt = f"{_label_modulo(m_fuerte)} ({modulos_activos[m_fuerte]:.2f} pts)"
        else:
            fuerte_txt = '-'
        if len(modulos_activos) > 1:
            m_debil = min(modulos_activos, key=modulos_activos.get)
            debil_txt = f"{_label_modulo(m_debil)} ({modulos_activos[m_debil]:.2f} pts)"
        elif len(modulos_activos) == 1:
            faltantes = [m for m in modulos_clave if m not in modulos_activos]
            debil_txt = ('Sin participación en: ' + ', '.join(_label_modulo(m) for m in faltantes)) if faltantes else '-'
        else:
            debil_txt = 'Sin operaciones registradas'

        filas.append([
            op.get('usuario') or '-',
            pct_ops,
            op.get('total_renglones') or 0,
            round(op.get('total_puntos') or 0, 2),
            round(op.get('puntos_picking') or 0, 2),
            round(op.get('puntos_chequeo') or 0, 2),
            round(op.get('puntos_inventario') or 0, 2),
            round(op.get('puntos_embalaje') or 0, 2),
            fuerte_txt,
            debil_txt,
        ])
    _escribir_tabla(
        ws, 3, headers, filas,
        anchos=[22, 14, 12, 14, 12, 12, 14, 13, 26, 30],
        filtro=True,
    )
    ws.sheet_view.showGridLines = False


def _hoja_detalle(wb, fecha_inicio, fecha_fin, incidencias):
    ws = wb.create_sheet('Detalle de Operaciones')
    ws['A1'] = f'Detalle de operaciones ({fecha_inicio} a {fecha_fin})'
    ws['A1'].font = _FONT_SUBTITULO

    headers = ['Fecha', 'Usuario', 'Módulo', 'Referencia', 'Renglones', 'Puntos']
    filas = []
    for inc in incidencias:
        filas.append([
            str(inc['fecha_registro'])[:10] if inc['fecha_registro'] else '-',
            inc['usuario'] if inc['usuario'] is not None else '-',
            str(inc['modulo']).capitalize() if inc['modulo'] else '-',
            inc['referencia_id'] if inc['referencia_id'] is not None else '-',
            inc['cantidad_renglones'] if inc['cantidad_renglones'] is not None else 0,
            round(inc['puntos_ganados'], 2) if inc['puntos_ganados'] is not None else 0.0,
        ])
    _escribir_tabla(ws, 3, headers, filas, anchos=[14, 20, 16, 22, 12, 12], filtro=True)
    ws.sheet_view.showGridLines = False


def generar_xlsx_reporte(incidencias, kpis, distribucion_modulo, rendimiento_operadores,
                          fecha_inicio, fecha_fin) -> io.BytesIO:
    wb = Workbook()
    _hoja_portada(wb, fecha_inicio, fecha_fin, kpis)
    _hoja_distribucion(wb, fecha_inicio, fecha_fin, distribucion_modulo)
    _hoja_operadores(wb, fecha_inicio, fecha_fin, rendimiento_operadores)
    _hoja_detalle(wb, fecha_inicio, fecha_fin, incidencias)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def register_xlsx_route(app):
    @app.route('/api/reporte/xlsx', methods=['GET', 'POST'], strict_slashes=False)
    def reporte_xlsx():
        if request.method == 'POST':
            params = request.get_json(silent=True) if request.is_json else request.form
            params = params or {}
        else:
            params = request.args

        fecha_inicio = params.get('fecha_inicio')
        fecha_fin = params.get('fecha_fin')

        if not fecha_fin:
            fecha_fin = datetime.now().strftime('%Y-%m-%d')
        if not fecha_inicio:
            fecha_inicio = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')

        try:
            datetime.strptime(fecha_inicio, '%Y-%m-%d')
            datetime.strptime(fecha_fin, '%Y-%m-%d')
        except (ValueError, TypeError):
            return jsonify({'error': 'Formato de fecha inválido. Use YYYY-MM-DD'}), 400

        try:
            incidencias = obtener_incidencias_rango(fecha_inicio, fecha_fin)
            kpis = obtener_resumen_kpis(fecha_inicio, fecha_fin)
            distribucion_modulo = obtener_distribucion_por_modulo(fecha_inicio, fecha_fin)
            rendimiento_operadores = obtener_rendimiento_operadores(fecha_inicio, fecha_fin)
        except Exception as e:
            return jsonify({'error': f'Error de base de datos: {e}'}), 500

        try:
            xlsx_buffer = generar_xlsx_reporte(
                incidencias, kpis, distribucion_modulo, rendimiento_operadores,
                fecha_inicio, fecha_fin,
            )
        except Exception as e:
            return jsonify({'error': f'Error generando XLSX: {e}'}), 500

        nombre = f'reporte_rendimiento_{fecha_inicio}_{fecha_fin}.xlsx'
        response = make_response(send_file(
            xlsx_buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=nombre,
        ))
        response.headers['Content-Disposition'] = f'attachment; filename="{nombre}"'
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        return response
